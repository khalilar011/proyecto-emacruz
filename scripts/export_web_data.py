"""Exporta la hoja Frecuencias del Excel original al JSON del prototipo web."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_FILE = PROJECT_ROOT / "data" / "Base_Emacruz_sectores_y_frecuencias.xlsx"
OUTPUT_FILE = PROJECT_ROOT / "web" / "data" / "frecuencias.json"
SHEET_NAME = "Frecuencias"

FIELD_MAP = {
    "ID": "id",
    "Sector publicado": "sector",
    "Zona": "zona",
    "Delimitación / área": "area",
    "Distritos indicados": "distritos",
    "Días publicados": "dias_publicados",
    "Feriados": "feriados",
    "Turno": "turno",
    "Inicio": "inicio",
    "Fin": "fin",
    "Nota de revisión": "nota_revision",
}

DAY_MAP = {
    "Lun": "lun",
    "Mar": "mar",
    "Mié": "mie",
    "Jue": "jue",
    "Vie": "vie",
    "Sáb": "sab",
    "Dom": "dom",
}


def clean_text(value: Any) -> str:
    """Conserva el dato publicado y normaliza únicamente vacíos reales."""

    if value is None:
        return "No indicado"
    text = str(value).strip()
    return text if text else "No indicado"


def day_is_active(value: Any, *, header: str, row_number: int) -> bool:
    """Convierte las marcas 1/0 del Excel sin inferir días adicionales."""

    if value is True or value == 1:
        return True
    if value is False or value in (0, None, ""):
        return False
    if isinstance(value, str):
        normalized = value.strip().lower()
        if normalized in {"1", "sí", "si", "x", "true"}:
            return True
        if normalized in {"0", "no", "false"}:
            return False
    raise ValueError(
        f"Marca de día inesperada en fila {row_number}, columna {header}: {value!r}"
    )


def export_frequencies() -> int:
    if not SOURCE_FILE.is_file():
        raise FileNotFoundError(f"No se encontró la fuente: {SOURCE_FILE}")

    workbook = load_workbook(SOURCE_FILE, read_only=True, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"No existe la hoja {SHEET_NAME!r} en {SOURCE_FILE.name}")

    worksheet = workbook[SHEET_NAME]
    rows = worksheet.iter_rows(values_only=True)
    headers = [clean_text(value) for value in next(rows)]
    header_positions = {header: position for position, header in enumerate(headers)}

    required_headers = [*FIELD_MAP, *DAY_MAP]
    missing_headers = [header for header in required_headers if header not in header_positions]
    if missing_headers:
        raise ValueError(f"Faltan columnas requeridas: {', '.join(missing_headers)}")

    records: list[dict[str, Any]] = []
    seen_ids: set[str] = set()

    for row_number, values in enumerate(rows, start=2):
        if not any(value is not None and str(value).strip() for value in values):
            continue

        record = {
            output_name: clean_text(values[header_positions[source_name]])
            for source_name, output_name in FIELD_MAP.items()
        }
        record["dias"] = {
            output_name: day_is_active(
                values[header_positions[source_name]],
                header=source_name,
                row_number=row_number,
            )
            for source_name, output_name in DAY_MAP.items()
        }
        record["fila_origen"] = row_number

        record_id = record["id"]
        if record_id == "No indicado":
            raise ValueError(f"La fila {row_number} no tiene ID")
        if record_id in seen_ids:
            raise ValueError(f"ID duplicado en la hoja Frecuencias: {record_id}")
        seen_ids.add(record_id)
        records.append(record)

    payload = {
        "fuente": SOURCE_FILE.name,
        "hoja": SHEET_NAME,
        "total_registros": len(records),
        "frecuencias": records,
    }

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_FILE.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    workbook.close()
    return len(records)


if __name__ == "__main__":
    total = export_frequencies()
    print(f"Exportados {total} registros a {OUTPUT_FILE.relative_to(PROJECT_ROOT)}")
